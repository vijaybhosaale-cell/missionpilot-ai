"""File exports: PDF (headless Chromium, correct Devanagari/Arabic shaping), Word
(.docx with language, right-to-left and script fonts set), MS Project XML (.xml),
Excel (.xlsx) and ZIP bundles."""
from __future__ import annotations

import asyncio
import base64
import datetime as dt
import io
import re
import zipfile
from html import escape
from xml.sax.saxutils import escape as xesc

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import i18n, render

# ---------------------------------------------------------------- Chromium (PDF + PNG)
_browser = None
_pw = None
_loop = None
_lock = None
pdf_engine_error: str | None = None


async def _get_browser():
    """One shared headless Chromium per event loop (started on first use)."""
    global _browser, _pw, _loop, _lock, pdf_engine_error
    loop = asyncio.get_running_loop()
    if _loop is not loop:
        _browser, _pw, _loop, _lock = None, None, loop, asyncio.Lock()
    async with _lock:
        if _browser is None or not _browser.is_connected():
            try:
                from playwright.async_api import async_playwright
                _pw = await async_playwright().start()
                _browser = await _pw.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
                pdf_engine_error = None
            except Exception as e:  # Chromium not installed (e.g. plain Windows run)
                pdf_engine_error = str(e)[:300]
                _browser = None
        return _browser


async def shutdown():
    global _browser, _pw
    if _browser:
        await _browser.close()
    if _pw:
        await _pw.stop()
    _browser = _pw = None


async def pdf_bytes(content: dict) -> bytes | None:
    """None when no Chromium is available; callers then fall back to the browser's own print-to-PDF."""
    b = await _get_browser()
    if not b:
        return None
    page = await b.new_page()
    try:
        await page.set_content(render.standalone(content, for_print=True), wait_until="load")
        loc = content["locale"]
        font = i18n.info(loc)["font_css"].replace('"', "'")
        head = (f'<div style="font-family:{font};font-size:8px;color:#6b7280;width:100%;padding:0 14mm;'
                f'display:flex;justify-content:space-between;direction:{content["dir"]}">'
                f'<span>{escape(content["project"])}</span><span>{escape(content["title"])}</span></div>')
        foot = (f'<div style="font-family:{font};font-size:8px;color:#6b7280;width:100%;padding:0 14mm;'
                f'display:flex;justify-content:space-between;direction:{content["dir"]}">'
                f'<span>MissionPilot AI · v{content.get("version", 1)} · {loc}</span>'
                f'<span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>')
        return await page.pdf(format="A4", print_background=True, display_header_footer=True, header_template=head,
                              footer_template=foot, margin={"top": "18mm", "bottom": "18mm", "left": "14mm", "right": "14mm"},
                              tagged=True)
    finally:
        await page.close()


async def svg_png(svg: str, locale: str, scale: float = 2.0) -> bytes | None:
    b = await _get_browser()
    if not b:
        return None
    m = re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg)
    w, h = (float(m.group(1)), float(m.group(2))) if m else (900, 400)
    page = await b.new_page(viewport={"width": int(w) + 2, "height": int(h) + 2}, device_scale_factor=scale)
    try:
        font = i18n.info(locale)["font_css"]
        await page.set_content(f'<html dir="{i18n.info(locale)["dir"]}"><body style="margin:0;background:#fff;font-family:{font}">{svg}</body></html>')
        el = await page.query_selector("svg")
        return await el.screenshot(type="png")
    finally:
        await page.close()


async def html_png(html: str, width: int = 1280, height: int = 800) -> bytes | None:
    b = await _get_browser()
    if not b:
        return None
    page = await b.new_page(viewport={"width": width, "height": height})
    try:
        await page.set_content(html)
        return await page.screenshot(full_page=True, type="png")
    finally:
        await page.close()


# ---------------------------------------------------------------- Word

def _set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _lang_run(run, loc):
    info = i18n.info(loc)
    sf = info["script_fonts"]
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), sf["docx_latin"])
    fonts.set(qn("w:hAnsi"), sf["docx_latin"])
    fonts.set(qn("w:cs"), sf["docx_cs"])
    if sf["docx_ea"]:
        fonts.set(qn("w:eastAsia"), sf["docx_ea"])
    if info["dir"] == "rtl":
        rtl = OxmlElement("w:rtl")
        rPr.append(rtl)
    lang = OxmlElement("w:lang")
    tag = info["docx_lang"]
    lang.set(qn("w:val"), tag if info["script"] == "latin" else "en-GB")
    if info["script"] in ("deva", "arab"):
        lang.set(qn("w:bidi"), tag)
    if info["script"] in ("jpan", "hans"):
        lang.set(qn("w:eastAsia"), tag)
    rPr.append(lang)


def _para_dir(p, loc):
    if i18n.info(loc)["dir"] == "rtl":
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.insert(0, bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _add_text(p, text, loc, bold=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
        szcs = OxmlElement("w:szCs")  # complex scripts (Devanagari, Arabic) use their own size
        szcs.set(qn("w:val"), str(int(size * 2)))
        r._r.get_or_add_rPr().append(szcs)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    _lang_run(r, loc)
    if bold:
        rPr = r._r.get_or_add_rPr()
        rPr.append(OxmlElement("w:bCs"))
    return r


def _table_rtl(table, loc):
    if i18n.info(loc)["dir"] == "rtl":
        tblPr = table._tbl.tblPr
        bv = OxmlElement("w:bidiVisual")
        tblPr.append(bv)


def _field(p, instr, loc):
    r = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)
    r.font.size = Pt(8)
    _lang_run(r, loc)



# Word is strict about the order of child elements; python-docx keeps it for what it adds itself,
# so the few elements added by hand (szCs, bCs, rtl, lang, bidi, bidiVisual) are put back in schema order.
_ORDER = {
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss",
            "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
            "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
            "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
            "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing",
              "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc", "pgBorders", "lnNumType",
               "pgNumType", "cols", "formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"],
}


def _fix_order(root):
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for tag, order in _ORDER.items():
        rank = {n: i for i, n in enumerate(order)}
        for el in root.iter(W + tag):
            kids = list(el)
            kids.sort(key=lambda k: rank.get(k.tag.replace(W, ""), 999))
            for k in kids:
                el.remove(k)
            for k in kids:
                el.append(k)


async def docx_bytes(content: dict) -> bytes:
    loc = content["locale"]
    info = i18n.info(loc)
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    if info["dir"] == "rtl":
        sectPr = sec._sectPr
        sectPr.append(OxmlElement("w:bidi"))
    # base style: fonts per script
    st = doc.styles["Normal"]
    st.font.size = Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    _szcs = OxmlElement("w:szCs"); _szcs.set(qn("w:val"), "21"); rpr.append(_szcs)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    sf = info["script_fonts"]
    rf.set(qn("w:ascii"), sf["docx_latin"]); rf.set(qn("w:hAnsi"), sf["docx_latin"]); rf.set(qn("w:cs"), sf["docx_cs"])
    if sf["docx_ea"]:
        rf.set(qn("w:eastAsia"), sf["docx_ea"])
    # core properties (metadata)
    cp = doc.core_properties
    cp.title = content["title"]
    cp.subject = content["project"]
    cp.language = info["docx_lang"]
    cp.version = f"v{content.get('version', 1)}"
    cp.author = "MissionPilot AI"
    cp.keywords = f"MissionPilot; {content['type']}; {loc}"
    cp.comments = f"project_id={content['project_id']}; language={loc}; version={content.get('version', 1)}"

    # header / footer
    hp = sec.header.paragraphs[0]
    _para_dir(hp, loc)
    _add_text(hp, f"{content['project']}  ·  {content['title']}", loc, size=8, color="6B7280")
    fp = sec.footer.paragraphs[0]
    _para_dir(fp, loc)
    _add_text(fp, f"MissionPilot AI · v{content.get('version', 1)} · {loc}    ", loc, size=8, color="6B7280")
    _field(fp, "PAGE", loc)
    _add_text(fp, " / ", loc, size=8, color="6B7280")
    _field(fp, "NUMPAGES", loc)

    kicker = i18n.t("doc.kicker_report" if content["kind"] == "report" else "doc.kicker_doc", loc)
    p = doc.add_paragraph(); _para_dir(p, loc); _add_text(p, kicker.upper(), loc, bold=True, size=9, color="4053D6")
    p = doc.add_paragraph(); _para_dir(p, loc); _add_text(p, content["title"], loc, bold=True, size=22)
    p = doc.add_paragraph(); _para_dir(p, loc); _add_text(p, content.get("subtitle") or content["project"], loc, size=13, color="6B7280")
    _kv_table(doc, content["meta"], loc)

    for b in content["blocks"]:
        t = b["t"]
        if t in ("h1", "h2"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if t == "h1" else 8)
            p.paragraph_format.keep_with_next = True
            _para_dir(p, loc)
            _add_text(p, b["text"], loc, bold=True, size=14 if t == "h1" else 12, color="4053D6" if t == "h1" else None)
        elif t == "p":
            p = doc.add_paragraph(); _para_dir(p, loc); _add_text(p, b["text"], loc)
        elif t == "note":
            p = doc.add_paragraph(); _para_dir(p, loc); _add_text(p, b["text"], loc, color="6B7280")
        elif t in ("bullets", "steps"):
            for i, item in enumerate(b["items"], 1):
                p = doc.add_paragraph(style="List Bullet" if t == "bullets" else "List Number")
                _para_dir(p, loc)
                _add_text(p, item, loc)
        elif t == "kv":
            _kv_table(doc, b["rows"], loc)
        elif t == "table":
            _data_table(doc, b, loc)
        elif t == "chart":
            png = await svg_png(b["svg"], loc)
            if png:
                p = doc.add_paragraph(); _para_dir(p, loc)
                p.add_run().add_picture(io.BytesIO(png), width=Cm(17.4))
            if b.get("caption"):
                p = doc.add_paragraph(); _para_dir(p, loc); _add_text(p, b["caption"], loc, size=9, color="6B7280")
    _fix_order(doc.element)
    for part in (sec.header, sec.footer):
        _fix_order(part._element)
    _fix_order(doc.styles.element)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _kv_table(doc, rows, loc):
    tb = doc.add_table(rows=0, cols=2)
    tb.style = "Table Grid"
    _table_rtl(tb, loc)
    for k, v in rows:
        cells = tb.add_row().cells
        _set_cell_bg(cells[0], "F5F6FB")
        for c, txt, bold in ((cells[0], k, True), (cells[1], str(v), False)):
            p = c.paragraphs[0]; _para_dir(p, loc); _add_text(p, txt, loc, bold=bold, size=9.5)
    doc.add_paragraph()


def _data_table(doc, b, loc):
    cols = b["cols"]
    tb = doc.add_table(rows=1, cols=len(cols))
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_rtl(tb, loc)
    num = set(b.get("num") or [])
    for i, c in enumerate(cols):
        cell = tb.rows[0].cells[i]
        _set_cell_bg(cell, "EEF0FF")
        p = cell.paragraphs[0]; _para_dir(p, loc); _add_text(p, c, loc, bold=True, size=9)
    # repeat header row on each page
    trPr = tb.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for r in b["rows"]:
        cells = tb.add_row().cells
        for i, v in enumerate(r):
            p = cells[i].paragraphs[0]
            _para_dir(p, loc)
            if i in num:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i18n.info(loc)["dir"] == "rtl" else WD_ALIGN_PARAGRAPH.RIGHT
            _add_text(p, str(v), loc, size=9)
    if b.get("widths"):
        total = sum(b["widths"])
        for row in tb.rows:
            for i, w in enumerate(b["widths"]):
                row.cells[i].width = Cm(17.4 * w / total)
    doc.add_paragraph()


# ---------------------------------------------------------------- MS Project XML (MSPDI)

def mspdi(data: dict, locale: str) -> bytes:
    """Microsoft Project XML: File > Open in MS Project (also ProjectLibre, GanttProject)."""
    p = data["project"]
    tasks = data["tasks"]
    uid = {t["id"]: i for i, t in enumerate(tasks, 1)}
    ruid = {m["id"]: i for i, m in enumerate(data["members"], 1)}
    phase_uid = {}
    out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
           '<Project xmlns="http://schemas.microsoft.com/project">',
           "<SaveVersion>14</SaveVersion>",
           f"<Name>{xesc(p['display_name'])}.xml</Name>", f"<Title>{xesc(p['display_name'])}</Title>",
           f"<Author>MissionPilot AI</Author>", f"<CreationDate>{dt.datetime.now().replace(microsecond=0).isoformat()}</CreationDate>",
           "<ScheduleFromStart>1</ScheduleFromStart>",
           f"<StartDate>{p['start_date']}T09:00:00</StartDate>", f"<FinishDate>{p['end_date']}T17:00:00</FinishDate>",
           "<CalendarUID>1</CalendarUID>", "<DefaultStartTime>09:00:00</DefaultStartTime>", "<DefaultFinishTime>17:00:00</DefaultFinishTime>",
           "<MinutesPerDay>480</MinutesPerDay>", "<MinutesPerWeek>2400</MinutesPerWeek>", "<DaysPerMonth>20</DaysPerMonth>",
           f"<CurrencyCode>{p['currency']}</CurrencyCode>", "<WeekStartDay>1</WeekStartDay>",
           "<Calendars><Calendar><UID>1</UID><Name>Standard</Name><IsBaseCalendar>1</IsBaseCalendar><WeekDays>"]
    for d in range(1, 8):
        working = 2 <= d <= 6
        out.append(f"<WeekDay><DayType>{d}</DayType><DayWorking>{1 if working else 0}</DayWorking>")
        if working:
            out.append("<WorkingTimes><WorkingTime><FromTime>09:00:00</FromTime><ToTime>13:00:00</ToTime></WorkingTime>"
                       "<WorkingTime><FromTime>14:00:00</FromTime><ToTime>18:00:00</ToTime></WorkingTime></WorkingTimes>")
        out.append("</WeekDay>")
    out.append("</WeekDays></Calendar></Calendars><Tasks>")
    out.append(f"<Task><UID>0</UID><ID>0</ID><Name>{xesc(p['display_name'])}</Name><OutlineLevel>0</OutlineLevel><Summary>1</Summary>"
               f"<Start>{p['start_date']}T09:00:00</Start><Finish>{p['end_date']}T18:00:00</Finish></Task>")
    n = 0
    next_uid = len(tasks) + 1
    for ph in data["phases"]:
        pts = [t for t in tasks if t["phase_id"] == ph["id"]]
        if not pts:
            continue
        n += 1
        phase_uid[ph["id"]] = next_uid
        s, e = min(t["start"] for t in pts), max(t["end"] for t in pts)
        work = [t for t in pts if not t["milestone"]]
        prog = round(sum(t["days"] * t["progress"] for t in work) / max(1, sum(t["days"] for t in work)))
        out.append(f"<Task><UID>{next_uid}</UID><ID>{n}</ID><Name>{xesc(ph['name'])}</Name><WBS>{ph['seq']}</WBS><OutlineNumber>{ph['seq']}</OutlineNumber>"
                   f"<OutlineLevel>1</OutlineLevel><Summary>1</Summary><Start>{s}T09:00:00</Start><Finish>{e}T18:00:00</Finish>"
                   f"<PercentComplete>{prog}</PercentComplete></Task>")
        next_uid += 1
        for t in pts:
            n += 1
            dur_h = 0 if t["milestone"] else t["days"] * 8
            finish_t = "09:00:00" if t["milestone"] else "18:00:00"
            notes = f"<Notes>{xesc(t['desc'])}</Notes>" if t["desc"] else ""
            preds = "".join(f"<PredecessorLink><PredecessorUID>{uid[q]}</PredecessorUID><Type>1</Type></PredecessorLink>"
                            for q in t["preds"] if q in uid)
            start_t = "18:00:00" if t["milestone"] else "09:00:00"
            out.append(
                f"<Task><UID>{uid[t['id']]}</UID><ID>{n}</ID><Name>{xesc(t['title'])}</Name><WBS>{t['wbs']}</WBS>"
                f"<OutlineNumber>{t['wbs']}</OutlineNumber><OutlineLevel>2</OutlineLevel><Type>1</Type>"
                f"<Start>{t['start']}T{start_t}</Start><Finish>{t['end']}T{finish_t if not t['milestone'] else '18:00:00'}</Finish>"
                f"<Duration>PT{dur_h}H0M0S</Duration><DurationFormat>7</DurationFormat><Work>PT{dur_h}H0M0S</Work>"
                f"<Milestone>{1 if t['milestone'] else 0}</Milestone><Summary>0</Summary>"
                f"<Priority>{ {'low': 300, 'medium': 500, 'high': 700, 'critical': 900}[t['priority']] }</Priority>"
                f"<PercentComplete>{t['progress']}</PercentComplete><PercentWorkComplete>{t['progress']}</PercentWorkComplete>"
                f"<ConstraintType>0</ConstraintType>{notes}{preds}</Task>")
    out.append("</Tasks><Resources><Resource><UID>0</UID><ID>0</ID><Type>1</Type><IsNull>0</IsNull></Resource>")
    for m in data["members"]:
        rate = m["rate_minor"] / 100 / 8
        out.append(f"<Resource><UID>{ruid[m['id']]}</UID><ID>{ruid[m['id']]}</ID><Name>{xesc(m['name'])}</Name><Type>1</Type>"
                   f"<Group>{xesc(m['role'])}</Group><MaxUnits>{m['capacity'] / 100:.2f}</MaxUnits>"
                   f"<StandardRate>{rate:.2f}</StandardRate><StandardRateFormat>2</StandardRateFormat><CalendarUID>1</CalendarUID></Resource>")
    out.append("</Resources><Assignments>")
    a = 0
    for t in tasks:
        if t["assignee_id"] in ruid:
            a += 1
            dur_h = 0 if t["milestone"] else t["days"] * 8
            out.append(f"<Assignment><UID>{a}</UID><TaskUID>{uid[t['id']]}</TaskUID><ResourceUID>{ruid[t['assignee_id']]}</ResourceUID>"
                       f"<Units>1</Units><Work>PT{dur_h}H0M0S</Work><PercentWorkComplete>{t['progress']}</PercentWorkComplete></Assignment>")
    out.append("</Assignments></Project>")
    return "\n".join(out).encode("utf-8")


# ---------------------------------------------------------------- Excel

def xlsx(data: dict, locale: str, what: str = "gantt") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    t = lambda k, **kw: i18n.t(k, locale, **kw)
    wb = Workbook()
    rtl = i18n.info(locale)["dir"] == "rtl"
    head = PatternFill("solid", fgColor="EEF0FF")
    bold = Font(bold=True)
    thin = Side(style="thin", color="E5E7EB")
    p = data["project"]
    cur = p["currency"]

    def sheet(title, cols, rows, widths=None, money_cols=()):
        ws = wb.create_sheet(title[:30])
        ws.sheet_view.rightToLeft = rtl
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font = head, bold
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in money_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for c in row:
                    c.number_format = f'#,##0.00 "{cur}"'
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    if what in ("gantt", "all"):
        ws = wb.create_sheet(t("gantt.title")[:30])
        ws.sheet_view.rightToLeft = rtl
        cols = [t("gantt.col_wbs"), t("gantt.col_task"), t("gantt.col_resource"), t("gantt.col_start"), t("gantt.col_end"),
                t("gantt.col_days"), t("gantt.col_pct")]
        tasks = data["tasks"]
        start = min(x["start"] for x in tasks)
        start -= dt.timedelta(days=start.weekday())
        end = max(x["end"] for x in tasks)
        weeks = []
        w = start
        while w <= end:
            weeks.append(w)
            w += dt.timedelta(days=7)
        ws.append(cols + [i18n.short_date(w, locale) for w in weeks])
        for c in ws[1]:
            c.fill, c.font = head, bold
            c.alignment = Alignment(text_rotation=90 if c.column > len(cols) else 0, horizontal="center")
        bar = PatternFill("solid", fgColor="C7CDFA")
        done = PatternFill("solid", fgColor="2F9E44")
        prog_fill = PatternFill("solid", fgColor="4053D6")
        ms = PatternFill("solid", fgColor="7048E8")
        for ph in data["phases"]:
            pts = [x for x in tasks if x["phase_id"] == ph["id"]]
            if not pts:
                continue
            ws.append([str(ph["seq"]), ph["name"]])
            for c in ws[ws.max_row]:
                c.font = bold
            for x in pts:
                ws.append([x["wbs"], x["title"], x["assignee"], x["start"], x["end"], x["days"], x["progress"] / 100])
                r = ws.max_row
                ws.cell(r, 4).number_format = ws.cell(r, 5).number_format = "yyyy-mm-dd"
                ws.cell(r, 7).number_format = "0%"
                for j, w in enumerate(weeks):
                    we = w + dt.timedelta(days=6)
                    if x["start"] <= we and x["end"] >= w:
                        cell = ws.cell(r, len(cols) + 1 + j)
                        if x["milestone"]:
                            cell.fill, cell.value = ms, "◆"
                        else:
                            span = (x["end"] - x["start"]).days + 1
                            done_until = x["start"] + dt.timedelta(days=span * x["progress"] / 100)
                            cell.fill = (done if x["progress"] >= 100 else prog_fill) if w < done_until else bar
                        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for i, wdt in enumerate([7, 46, 18, 12, 12, 7, 7], 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        for j in range(len(weeks)):
            ws.column_dimensions[get_column_letter(len(cols) + 1 + j)].width = 3.2
        ws.freeze_panes = ws.cell(2, len(cols) + 1)
    if what in ("tasks", "all", "gantt"):
        sheet(t("nav.tasks"), [t("gantt.col_wbs"), t("gantt.col_task"), t("doc.c.phase"), t("gantt.col_resource"), t("gantt.col_start"),
                               t("gantt.col_end"), t("gantt.col_days"), t("gantt.col_pct"), t("doc.c.status"), t("doc.c.priority"),
                               t("progress.done_days"), t("progress.left_days")],
              [[x["wbs"], x["title"], x["phase"], x["assignee"], x["start"], x["end"], x["days"], x["progress"],
                t(f"status.{x['status']}"), t(f"priority.{x['priority']}"), x["done_days"], x["left_days"]] for x in data["tasks"]],
              [7, 46, 18, 18, 12, 12, 7, 7, 14, 10, 10, 10])
    if what in ("budget", "all"):
        sheet(t("nav.budget"), [t("doc.c.category"), t("doc.c.planned"), t("doc.c.actual"), t("doc.c.variance")],
              [[b["label"], b["planned"] / 100, b["actual"] / 100, b["variance"] / 100] for b in data["budget"]],
              [30, 16, 16, 16], money_cols=(2, 3, 4))
    if what in ("risks", "all"):
        sheet(t("nav.risks"), [t("doc.c.id"), t("doc.c.risk"), t("doc.c.category"), t("doc.c.p"), t("doc.c.i"), t("doc.c.score"),
                               t("doc.c.owner"), t("doc.c.mitigation"), t("doc.c.status")],
              [[r["code"], r["title"], t(f"riskcat.{r['category']}"), r["probability"], r["impact"], r["score"], r["owner"],
                r["mitigation"], t(f"riskstatus.{r['status']}")] for r in data["risks"]],
              [7, 44, 14, 5, 5, 7, 18, 60, 12])
    if what in ("team", "all"):
        from . import analytics
        wl = analytics.workload(data)
        cols = [t("doc.c.member"), t("doc.c.role")] + [i18n.short_date(w, locale) for w in wl["weeks"]]
        sheet(t("nav.team"), cols, [[r["member"]["name"], r["member"]["role"]] + r["hours"] for r in wl["rows"]],
              [22, 22] + [8] * len(wl["weeks"]))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def zip_bytes(files: list[tuple[str, bytes]]) -> bytes:
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    return bio.getvalue()


def safe_name(s: str) -> str:
    s = re.sub(r"[^\w\-]+", "-", s, flags=re.U).strip("-")
    return s[:60] or "project"
