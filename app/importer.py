"""Import existing projects from Excel/CSV, Trello (JSON export) and Jira (CSV export).

Step 1 parse the file into columns + rows (nothing saved in the project yet),
step 2 the user maps columns to fields (guessed automatically),
step 3 validate and preview (errors shown per row), step 4 create the project.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re

from . import db, i18n, library as L, ml, planner

FIELDS = ["title", "phase", "assignee", "start", "end", "progress", "status", "priority"]
REQUIRED = ["title"]
SYNONYMS = {
    "title": ["title", "task", "task name", "name", "summary", "card", "card name", "subject", "tâche", "nom", "titre", "intitulé",
              "résumé", "काम", "टास्क", "नाव"],
    "phase": ["phase", "list", "list name", "epic", "epic link", "group", "stage", "sprint", "issue type", "étape", "liste", "groupe", "टप्पा"],
    "assignee": ["assignee", "owner", "assigned to", "members", "member", "responsible", "resource names", "resource", "responsable",
                 "assigné", "assigné à", "membre", "जबाबदार", "मालक"],
    "start": ["start", "start date", "begin", "created", "début", "date de début", "सुरुवात"],
    "end": ["end", "end date", "due", "due date", "finish", "deadline", "fin", "date de fin", "échéance", "date d'échéance", "शेवट", "मुदत"],
    "progress": ["progress", "% complete", "percent complete", "progress %", "done %", "avancement", "% achevé", "प्रगती"],
    "status": ["status", "state", "statut", "état", "स्थिती", "resolution"],
    "priority": ["priority", "priorité", "प्राधान्य", "importance"],
}
DONE = {"done", "closed", "resolved", "complete", "completed", "finished", "terminé", "termine", "fait", "clos", "fermé", "पूर्ण", "true", "yes"}
DOING = {"in progress", "doing", "in review", "review", "en cours", "en revue", "चालू", "सुरू"}
BLOCKED = {"blocked", "bloqué", "bloque", "on hold", "impediment", "अडकले"}
PRIO = {"highest": "critical", "blocker": "critical", "critical": "critical", "critique": "critical", "urgent": "critical",
        "high": "high", "haute": "high", "élevée": "high", "major": "high", "medium": "medium", "moyenne": "medium", "normal": "medium",
        "low": "low", "lowest": "low", "basse": "low", "minor": "low", "trivial": "low"}


# ---------------------------------------------------------------- parsing

def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").replace("﻿", "")).strip()


def parse_file(filename: str, raw: bytes) -> dict:
    """-> {source, columns, rows, name}. Raises ValueError with a translation key on bad input."""
    name = (filename or "").lower()
    if len(raw) > 5_000_000:
        raise ValueError("imp.err.too_big")
    if name.endswith(".json"):
        return _trello(raw)
    if name.endswith(".xlsx"):
        return _xlsx(raw)
    if name.endswith(".csv") or name.endswith(".txt"):
        return _csv(raw)
    raise ValueError("imp.err.type")


def _csv(raw: bytes) -> dict:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise ValueError("imp.err.empty")
    header = [_norm_header(h) for h in rows[0]]
    # Jira repeats some column names (e.g. several "Sprint"): keep the first
    seen, cols = set(), []
    for i, h in enumerate(header):
        cols.append(h if h not in seen else f"{h} ({i})")
        seen.add(h)
    data = [{cols[i]: (r[i].strip() if i < len(r) else "") for i in range(len(cols))} for r in rows[1:]][:2000]
    source = "jira" if {"Issue key", "Summary"} & set(header) else "csv"
    return dict(source=source, columns=cols, rows=data, name="")


def _xlsx(raw: bytes) -> dict:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError("imp.err.read")
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True) if r and any(c not in (None, "") for c in r)]
    if len(rows) < 2:
        raise ValueError("imp.err.empty")
    cols = [_norm_header(h) or f"Column {i + 1}" for i, h in enumerate(rows[0])]

    def cell(v):
        if isinstance(v, dt.datetime):
            return v.date().isoformat()
        if isinstance(v, dt.date):
            return v.isoformat()
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return "" if v is None else str(v).strip()
    data = [{cols[i]: cell(r[i]) if i < len(r) else "" for i in range(len(cols))} for r in rows[1:]][:2000]
    return dict(source="excel", columns=cols, rows=data, name=ws.title if ws.title not in ("Sheet", "Sheet1", "Feuil1") else "")


def _trello(raw: bytes) -> dict:
    try:
        b = json.loads(raw.decode("utf-8-sig"))
    except Exception:
        raise ValueError("imp.err.read")
    if not isinstance(b, dict) or "cards" not in b:
        raise ValueError("imp.err.trello")
    lists = {x.get("id"): x.get("name", "") for x in b.get("lists", [])}
    members = {x.get("id"): x.get("fullName") or x.get("username", "") for x in b.get("members", [])}
    cols = ["Name", "List", "Members", "Start", "Due", "Done", "Labels"]
    rows = []
    for c in b["cards"]:
        if c.get("closed"):
            continue
        rows.append({"Name": c.get("name", ""), "List": lists.get(c.get("idList"), ""),
                     "Members": ", ".join(members.get(m, "") for m in c.get("idMembers", []) if members.get(m)),
                     "Start": (c.get("start") or "")[:10], "Due": (c.get("due") or "")[:10],
                     "Done": _trello_state(c, lists.get(c.get("idList"), "")), "Labels": ", ".join(lb.get("name", "") for lb in c.get("labels", []))})
    if not rows:
        raise ValueError("imp.err.empty")
    return dict(source="trello", columns=cols, rows=rows[:2000], name=b.get("name", ""))


def _trello_state(card: dict, list_name: str) -> str:
    ln = list_name.strip().lower()
    if card.get("dueComplete") or ln in DONE:
        return "done"
    if ln in DOING:
        return "in progress"
    if ln in BLOCKED:
        return "blocked"
    return ""


def guess_mapping(columns: list[str], source: str = "") -> dict:
    out = {}
    lowcols = {c: c.lower().strip() for c in columns}
    for f in FIELDS:
        for syn in SYNONYMS[f]:
            hit = next((c for c, lc in lowcols.items() if lc == syn and c not in out.values()), None)
            if hit:
                out[f] = hit
                break
    if source == "trello":
        out["status"] = "Done"
    return out


# ---------------------------------------------------------------- validation

def parse_date(v: str):
    v = (v or "").strip()
    if not v:
        return None
    v = v.translate(str.maketrans("०१२३४५६७८९", "0123456789"))
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d.%m.%Y", "%d-%m-%Y", "%Y/%m/%d", "%d/%b/%y %I:%M %p", "%d/%b/%y",
                "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d %b %Y", "%b %d, %Y"):
        try:
            return dt.datetime.strptime(v[:len(v)], fmt).date()
        except ValueError:
            continue
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", v)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3]))
    raise ValueError(v)


def _progress(row: dict, mp: dict) -> tuple[int, str]:
    status_raw = row.get(mp.get("status", ""), "").strip().lower() if mp.get("status") else ""
    prog = None
    if mp.get("progress"):
        pv = row.get(mp["progress"], "").replace("%", "").replace(",", ".").strip()
        if pv:
            p = float(pv)
            if 0 < p <= 1 and "." in pv:
                p *= 100
            prog = int(round(max(0, min(100, p))))
    blocked = status_raw in BLOCKED
    if prog is None:
        if status_raw in DONE:
            prog = 100
        elif status_raw in DOING:
            prog = 50
        else:
            prog = 0
    status = "done" if prog >= 100 else ("blocked" if blocked else ("in_progress" if prog > 0 else "todo"))
    return prog, status


def validate(rows: list[dict], mp: dict, today: dt.date) -> tuple[list[dict], list[tuple[int, str, str]]]:
    items, errors = [], []
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        title = row.get(mp.get("title", ""), "").strip() if mp.get("title") else ""
        if not title:
            errors.append((n, "imp.err.no_title", ""))
            continue
        try:
            start = parse_date(row.get(mp["start"], "")) if mp.get("start") else None
        except ValueError as e:
            errors.append((n, "imp.err.bad_date", str(e)))
            continue
        try:
            end = parse_date(row.get(mp["end"], "")) if mp.get("end") else None
        except ValueError as e:
            errors.append((n, "imp.err.bad_date", str(e)))
            continue
        if start and end and end < start:
            errors.append((n, "imp.err.end_before_start", f"{start} > {end}"))
            continue
        try:
            prog, status = _progress(row, mp)
        except ValueError:
            errors.append((n, "imp.err.bad_progress", row.get(mp.get("progress", ""), "")))
            continue
        assignee = row.get(mp.get("assignee", ""), "").split(",")[0].strip() if mp.get("assignee") else ""
        prio = PRIO.get(row.get(mp.get("priority", ""), "").strip().lower(), "medium") if mp.get("priority") else "medium"
        group = row.get(mp.get("phase", ""), "").strip() if mp.get("phase") else ""
        items.append(dict(row=n, title=title[:200], group=group[:80], assignee=assignee[:80], start=start, end=end,
                          progress=prog, status=status, priority=prio))
    # fill missing dates: 5 working days, chained per assignee
    cursor = {}
    for it in items:
        if it["start"] and not it["end"]:
            it["end"] = planner.add_workdays(it["start"], 4)
        elif it["end"] and not it["start"]:
            s = it["end"]
            for _ in range(4):
                s -= dt.timedelta(days=1)
                while not planner.is_workday(s):
                    s -= dt.timedelta(days=1)
            it["start"] = s
        elif not it["start"] and not it["end"]:
            s = cursor.get(it["assignee"], planner.next_workday(today))
            it["start"], it["end"] = s, planner.add_workdays(s, 4)
            it["dates_guessed"] = True
        cursor[it["assignee"]] = planner.add_workdays(it["end"], 1)
        it["start"] = planner.next_workday(it["start"])
        if it["end"] < it["start"]:
            it["end"] = it["start"]
    return items, errors


# ---------------------------------------------------------------- create the project

def phase_for(group: str, title: str) -> str:
    text = f"{group} {title}".strip()
    return ml.suggest_phase(text, "build")


def create(owner_id: int, name: str, items: list[dict], *, budget: float, currency: str, locale: str) -> int:
    from .planner import new_project_id
    if not items:
        raise ValueError("imp.err.empty")
    ptype = planner.detect_type(name + " " + " ".join(i["title"] for i in items[:30]))
    start = min(i["start"] for i in items)
    end = max(i["end"] for i in items)
    now = db.now()
    pid = new_project_id()
    db.ex("""INSERT INTO projects(id,owner_id,name,description,project_type,canonical_locale,currency,budget_minor,team_size,start_date,end_date,
             engine,answers_json,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
          (pid, owner_id, name[:120] or "Imported project", "", ptype, locale, currency, planner.to_minor(budget, currency),
           0, start.isoformat(), end.isoformat(), "import", json.dumps(dict(source="import"), ensure_ascii=False), now))
    names = []
    for i in items:
        if i["assignee"] and i["assignee"] not in names:
            names.append(i["assignee"])
    names = names[:40] or [""]
    mid = {}
    for k, nm in enumerate(names, 1):
        role = "project_manager" if k == 1 else "developer"
        mid[nm] = db.ex("INSERT INTO members(project_id,seq,name,role_key,daily_rate_minor) VALUES(?,?,?,?,?)",
                        (pid, k, nm, role, planner.to_minor(L.DAILY_RATE_EUR["dev"], currency)))
    db.ex("UPDATE projects SET team_size=? WHERE id=?", (len(names), pid))
    phase_id = {}
    for k, code in enumerate(L.PHASES, 1):
        phase_id[code] = db.ex("INSERT INTO phases(project_id,seq,code,name_key,name_text) VALUES(?,?,?,?,?)",
                               (pid, k, code, f"phase.{code}", ""))
    phase_seq = {c: k for k, c in enumerate(L.PHASES, 1)}
    for it in items:
        it["phase"] = phase_for(it["group"], it["title"])
    counters: dict = {}
    for seq, it in enumerate(sorted(items, key=lambda x: (phase_seq[x["phase"]], x["start"], x["row"])), 1):
        counters[it["phase"]] = counters.get(it["phase"], 0) + 1
        days = max(1, planner.workdays_between(it["start"], it["end"]))
        tid = db.ex("""INSERT INTO tasks(project_id,phase_id,seq,wbs,code,title_text,start_date,end_date,duration_days,progress,status,
                       priority,assignee_id,is_milestone,effort_hours,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,0,?,?)""",
                    (pid, phase_id[it["phase"]], seq, f"{phase_seq[it['phase']]}.{counters[it['phase']]}", f"imp{seq}", it["title"],
                     it["start"].isoformat(), it["end"].isoformat(), days, it["progress"], it["status"], it["priority"],
                     mid.get(it["assignee"] or names[0]), days * 8, now))
        if it["progress"]:
            db.ex("INSERT INTO activity(project_id,user_id,at,kind,task_id,old_value,new_value) VALUES(?,?,?,?,?,?,?)",
                  (pid, owner_id, now, "progress", tid, "0", str(it["progress"])))
    risks = L.COMMON_RISKS + L.TYPE_RISKS.get(ptype, [])
    for k, (code, cat, p, i, _o) in enumerate(risks, 1):
        db.ex("""INSERT INTO risks(project_id,seq,code,title_key,mitigation_key,category,probability,impact,owner_id)
                 VALUES(?,?,?,?,?,?,?,?,?)""", (pid, k, f"R{k:02d}", f"risk.{code}.t", f"risk.{code}.m", cat, p, i, mid[names[0]]))
    split = L.BUDGET_SPLIT.get(ptype, L.BUDGET_SPLIT["generic"])
    total = planner.to_minor(budget, currency)
    acc = 0
    for k, (cat, pct) in enumerate(zip(L.BUDGET_CATEGORIES, split), 1):
        amt = int(round(total * pct / 100)) if k < len(split) else total - acc
        acc += amt
        db.ex("INSERT INTO budget_lines(project_id,seq,category,planned_minor) VALUES(?,?,?,?)", (pid, k, cat, amt))
    db.ex("INSERT INTO activity(project_id,user_id,at,kind,new_value) VALUES(?,?,?,?,?)", (pid, owner_id, now, "created", "import"))
    return pid


# ---------------------------------------------------------------- templates

TEMPLATE_ROWS = [
    ["Kick-off meeting", "Initiation", "Anaya Patil", "2026-11-02", "2026-11-02", "100", "Done", "High"],
    ["Gather requirements", "Planning", "Lucas Martin", "2026-11-03", "2026-11-10", "50", "In progress", "High"],
    ["Design the architecture", "Design", "Sofia Rossi", "2026-11-11", "2026-11-17", "0", "To do", "Medium"],
    ["Build the login page", "Build", "Lucas Martin", "2026-11-18", "2026-11-27", "0", "To do", "Medium"],
    ["User acceptance tests", "Testing", "Mei Chen", "2026-11-30", "2026-12-04", "0", "To do", "High"],
]
TEMPLATE_HEADER = ["Task", "Phase", "Owner", "Start", "End", "Progress %", "Status", "Priority"]


def template_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TEMPLATE_HEADER)
    w.writerows(TEMPLATE_ROWS)
    return buf.getvalue().encode("utf-8-sig")


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    ws.append(TEMPLATE_HEADER)
    for r in TEMPLATE_ROWS:
        ws.append([dt.date.fromisoformat(v) if k in (3, 4) else (int(v) if k == 5 else v) for k, v in enumerate(r)])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4053D6")
    for col, w in zip("ABCDEFGH", (34, 14, 18, 12, 12, 12, 14, 10)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def label(field: str, loc: str) -> str:
    return i18n.t(f"imp.f.{field}", loc)
